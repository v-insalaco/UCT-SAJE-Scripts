# Script to report on the 6 questions required for the SAJE Audit
# Working to get good results for all questions
# Last question may need some eyeballing to infer page updates at current as it shows last editor

import requests, json, os, glob
import pandas as pd
from tqdm import tqdm
from bs4 import BeautifulSoup
from datetime import datetime
import pytz
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill
start = datetime.now()

with open(os.path.expanduser("~") + r'/config.json') as json_data_file:
    configuration = json.load(json_data_file)
access_token = configuration["canvas"]["access_token"]
baseUrl = "https://" + configuration["canvas"]["host"] + "/api/v1/"
header = {'Authorization': 'Bearer ' + access_token}
csv_file = glob.glob('*.csv')
if len(csv_file) != 1:
    raise ValueError('should be only one file in the current directory')
csvfilename = csv_file[0]
if not os.path.exists(csvfilename):
    raise FileNotFoundError(f"CSV file '{csvfilename}' not found.")
pd.options.mode.chained_assignment = None

df_courses = pd.read_csv(csvfilename, encoding='utf-8')
df_courses.columns = [c.strip().lower().replace(" ", "_") for c in df_courses.columns]
if 'courseid' in df_courses.columns:
    df_courses.rename(columns={"courseid": "course_id"}, inplace=True)

if 'course_id' not in df_courses.columns:
    raise KeyError("CSV must contain a 'course_id' column")

output_file = os.path.join(os.getcwd(), "assessment_overview_audit.xlsx")

# -------------------------
def safe_get(url, params=None):
    try:
        r = requests.get(url, headers=header, params=params)
        if r.status_code == 200:
            return r.json()
        return None
    except Exception as e:
        print(f"Warning: failed GET {url} -> {e}")
        return None

def paginate(url, params=None):
    results = []
    try:
        r = requests.get(url, headers=header, params=params)
        if r.status_code != 200:
            return []
        results.extend(r.json())
        while r.links.get("next"):
            r = requests.get(r.links["next"]["url"], headers=header)
            if r.status_code == 200:
                results.extend(r.json())
            else:
                break
    except Exception:
        pass
    return results

# -------------------------
# Course-level queries
# -------------------------
def get_course_info(course_id):
    url = f"{baseUrl}courses/{course_id}"
    data = safe_get(url)
    if not data:
        return None

    # Get the account_id first
    account_id = data.get("account_id")
    sis_account_id = "N/A"

    # If we have an account_id, make a call to get sis_account_id
    if account_id:
        account_url = f"{baseUrl}accounts/{account_id}"
        account_data = safe_get(account_url)
        if account_data and "sis_account_id" in account_data:
            sis_account_id = account_data["sis_account_id"]

    return {
        "Course_ID": course_id,
        "Course_url": f"https://{configuration["canvas"]["host"]}/courses/{course_id}",
        "Long_name": data.get("name", ""),
        "Sub_account": sis_account_id,
        "Course_Published": "Yes" if data.get("workflow_state") == "available" else "No"
    }

def get_module_assessment_page(course_id):
    slug = "module-assessment-overview"
    url = f"{baseUrl}courses/{course_id}/pages/{slug}"
    page = safe_get(url)
    if page:
        return page
    # fallback: search all pages
    pages = paginate(f"{baseUrl}courses/{course_id}/pages")
    for p in pages:
        if p.get("title", "").strip().lower() == "module assessment overview":
            return p
    return None

def check_homepage_link(course_id):

    """Determines the homepage type (wiki or syllabus) for a course and searches
    for a 'Module Assessment Overview' link pointing to the correct page."""

    try:
        # 1. Identify homepage type
        course_url = f"{baseUrl}courses/{course_id}"
        course_data = safe_get(course_url)
        if not course_data:
            print(f"[{course_id}] Failed to fetch course data")
            return "No"

        default_view = course_data.get("default_view", "unknown")
        print(f"[{course_id}] Homepage type: {default_view}")

        # 2. Fetch the homepage HTML depending on type
        html = ""
        if default_view == "wiki":
            front_page_url = f"{baseUrl}courses/{course_id}/front_page"
            page = safe_get(front_page_url)
            if page and page.get("body"):
                html = page["body"]
            else:
                print(f"[{course_id}] No front page found")
                return "No"

        elif default_view == "syllabus":
            syllabus_url = f"{baseUrl}courses/{course_id}?include[]=syllabus_body"
            syllabus = safe_get(syllabus_url)
            if syllabus and syllabus.get("syllabus_body"):
                html = syllabus["syllabus_body"]
            else:
                print(f"[{course_id}] No syllabus body found")
                return "No"

        else:
            # e.g. modules or assignments page – not an HTML homepage
            print(f"[{course_id}] Non-HTML homepage type ({default_view}), skipping.")
            return "No"

        if not html.strip():
            print(f"[{course_id}] Empty homepage HTML")
            return "No"
        
        # 3. Parse and look for the target link
        soup = BeautifulSoup(html, "html.parser")
        for a in soup.find_all("a", href=True):
            href = a["href"].lower()
            if ("module-assessment-overview" in href):
                print(f"[{course_id}] Found valid link: {href}")
                return "Yes"

        print(f"[{course_id}] No matching Module Assessment Overview link found.")
        
        return "No"

    except Exception as e:
        print(f"[{course_id}] Exception in check_homepage_link: {e}")
        
        return "No"
    

def check_module_block(course_id, page):
    mods = paginate(f"{baseUrl}courses/{course_id}/modules")
    if not mods:
        return "No module"
    for m in mods:
        if m.get("name", "").strip().lower() == "module assessment overview":
            published = "Published" if m.get("published") else "Unpublished"
            # check items
            items = paginate(f"{baseUrl}courses/{course_id}/modules/{m['id']}/items")
            found = False
            if page:
                for it in items:
                    if str(page.get("url", "")).endswith(it.get("page_url", "")):
                        found = True
                        break
            return ("Yes" if found else "No") + f" - {published}"
    return "No module"

def get_page_updated(page):
    if not page:
        return "N/A"

    author = ""
    if page.get("last_edited_by"):
        author = page["last_edited_by"].get("display_name", "")

    updated_raw = page.get("updated_at", "")
    if not updated_raw:
        return author if author else "N/A"

    try:
        # Canvas timestamps are UTC, e.g. 2025-10-03T10:55:12Z
        utc_dt = datetime.strptime(updated_raw, "%Y-%m-%dT%H:%M:%SZ")
        london_tz = pytz.timezone("Europe/London")
        london_dt = pytz.utc.localize(utc_dt).astimezone(london_tz)
        formatted_time = london_dt.strftime("%d/%m/%y %H:%M:%S")
        return f"{author} - {formatted_time}" if author else formatted_time
    except Exception:
        return author if author else "N/A"

# -------------------------
# Main
# -------------------------
def main():

    results = []

    with tqdm(total=len(df_courses)) as pbar:
        for _, row in df_courses.iterrows():
            course_id = str(row['course_id']).strip()
            
            print('\n')
            print("Auditing:", course_id)

            info = get_course_info(course_id)
            if not info:
                pbar.update(1)
                continue

            page = get_module_assessment_page(course_id)
            info["Page_Available"] = "Yes" if page else "No"
            info["Homepage_Link"] = check_homepage_link(course_id)
            info["Page_Published"] = page.get("published", False) if page else "N/A"
            info["Module_Published"] = check_module_block(course_id, page)
            info["Page_Updated"] = get_page_updated(page)

            results.append(info)
            pbar.update(1)

    df = pd.DataFrame(results)
    df.to_excel(output_file, index=False)

    # -------------------------
    # Format Excel
    # -------------------------
    wb = load_workbook(output_file)
    ws = wb.active

    # Autofit
    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 2

    # Highlight No/Unpublished
    red_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
    header = [c.value for c in ws[1]]
    for target_col in ["Course_Published", "Page_Available", "Homepage_Link", "Page_Published", "Module_Published"]:
        if target_col in header:
            col_idx = header.index(target_col) + 1

            for col in ws.iter_cols(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in col:

                    if cell.value and str(cell.value).lower().startswith("no"):
                        cell.fill = red_fill
                    if str(cell.value).lower() == "unpublished":
                        cell.fill = red_fill

    wb.save(output_file)
    print("Excel file created")

# -------------------------

if __name__ == "__main__": main()
print("Running time:", datetime.now() - start)